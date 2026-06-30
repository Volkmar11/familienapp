# -*- coding: utf-8 -*-
"""
Erzeugt eine Microsoft-Forms-taugliche Umfrage als Word-Dokument (.docx).
Das Dokument kann in Microsoft Forms über "Neues Formular" > "Importieren"
(bzw. Copilot-/KI-Import aus Word) hochgeladen oder als Vorlage zum
manuellen Anlegen genutzt werden.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# ---------- Basis-Styles ----------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = ACCENT
    return p


def section(text):
    p = doc.add_paragraph()
    p.space_before = Pt(12)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = ACCENT
    return p


def intro(text):
    p = doc.add_paragraph(text)
    return p


def question(num, text, qtype, required):
    p = doc.add_paragraph()
    r = p.add_run(f"{num}. {text}")
    r.bold = True
    pflicht = "Pflichtfrage" if required else "freiwillig"
    meta = doc.add_paragraph()
    rm = meta.add_run(f"[{qtype} · {pflicht}]")
    rm.italic = True
    rm.font.size = Pt(9)
    rm.font.color.rgb = RGBColor(0x70, 0x70, 0x70)


def options(opts):
    for o in opts:
        doc.add_paragraph(o, style="List Bullet")


def matrix(rows, scale):
    p = doc.add_paragraph()
    r = p.add_run("Skala je Zeile: " + " – ".join(scale))
    r.italic = True
    r.font.size = Pt(9)
    for row in rows:
        doc.add_paragraph(row, style="List Bullet")


# ================= TITEL & EINLEITUNG =================
h1("Künstliche Intelligenz an unserer Schule")
sub = doc.add_paragraph()
rs = sub.add_run("Befragung des Kollegiums der Engelbert-Bohn-Schule Karlsruhe")
rs.italic = True
rs.font.size = Pt(12)

intro(
    "Liebe Kolleginnen und Kollegen,\n\n"
    "im kommenden Schuljahr möchten wir den Einsatz von Künstlicher Intelligenz (KI) "
    "an unserer Schule schrittweise und gut überlegt einführen. Bevor wir konkrete "
    "Entscheidungen treffen, ist uns Ihre Einschätzung wichtig: Welche Erfahrungen "
    "bestehen bereits, welche Wünsche und welche Bedenken gibt es im Kollegium?\n\n"
    "Diese Befragung ist anonym. Angaben zu Fachbereich, Funktion oder Name am Ende "
    "sind ausdrücklich freiwillig. Es gibt keine richtigen oder falschen Antworten – "
    "auch eine skeptische Haltung ist eine wertvolle Rückmeldung.\n\n"
    "Die Bearbeitung dauert etwa 10–15 Minuten. Vielen Dank, dass Sie sich die Zeit nehmen!"
)

doc.add_paragraph()

# ================= ABSCHNITT A =================
section("Abschnitt A – Bekanntheit und bisherige Erfahrung")

question(1,
         "Wie gut kennen Sie die folgenden KI-Tools und Begriffe? Bitte schätzen Sie für jede Zeile ein.",
         "Matrix / Likert (Einfachauswahl je Zeile)", True)
matrix(
    ["ChatGPT", "Claude", "Google Gemini", "Microsoft Copilot", "Perplexity",
     "Fobizz", "F13 (KI-Assistent BW)", "schulKI", "Eduhu",
     "Large Language Models (LLM)", "KI-Chatbots", "KI-Agenten", "Prompting (Eingaben formulieren)"],
    ["Kenne ich nicht", "Schon gehört", "Schon ausprobiert", "Nutze ich regelmäßig"]
)

question(2, "Wie häufig nutzen Sie KI-Tools privat (außerhalb der Schule)?",
         "Einfachauswahl", True)
options(["Nie", "Selten (einige Male im Jahr)", "Gelegentlich (monatlich)",
         "Regelmäßig (wöchentlich)", "Täglich oder fast täglich"])

question(3, "Welche KI-Tools nutzen Sie – privat oder beruflich – bereits? (Mehrfachauswahl)",
         "Mehrfachauswahl", True)
options(["ChatGPT", "Claude", "Google Gemini", "Microsoft Copilot", "Perplexity",
         "Fobizz", "F13", "schulKI", "Eduhu", "Andere (bitte angeben)",
         "Ich nutze bisher keine KI-Tools"])

question(4, "Wie greifen Sie überwiegend auf KI-Tools zu? (Mehrfachauswahl)",
         "Mehrfachauswahl", False)
options(["Über den Browser (z. B. am PC/Laptop)", "Über eine App auf dem Smartphone/Tablet",
         "Über eine Desktop-Anwendung", "Bisher gar nicht"])

question(5, "Welche Version nutzen Sie überwiegend?",
         "Einfachauswahl", True)
options(["Nur kostenlose Versionen", "Mindestens ein kostenpflichtiges Abo (privat bezahlt)",
         "Kostenpflichtiges Abo, über andere/dienstlich bereitgestellt",
         "Ich nutze bisher keine KI-Tools"])

# ================= ABSCHNITT B =================
section("Abschnitt B – Berufliche Nutzung von KI")

question(6, "Wie häufig nutzen Sie KI für schulische bzw. berufliche Aufgaben?",
         "Einfachauswahl", True)
options(["Nie", "Selten", "Gelegentlich (monatlich)", "Regelmäßig (wöchentlich)",
         "Täglich oder fast täglich"])

question(7, "Für welche schulischen Aufgaben haben Sie KI bereits genutzt? (Mehrfachauswahl)",
         "Mehrfachauswahl", False)
options(["Unterrichtsvorbereitung / Stundenplanung", "Erstellung von Arbeitsblättern",
         "Erstellung von Klassenarbeiten / Aufgaben", "Erstellung von Musterlösungen",
         "Differenzierung von Materialien", "Formulierung von Elternbriefen / Schülerhinweisen",
         "Recherche", "Zusammenfassung von Texten", "Ideenfindung / Brainstorming",
         "Korrekturunterstützung / Feedback", "Erstellung von Präsentationen",
         "Erstellung von Moodle-Inhalten", "Verwaltungs- und Organisationsaufgaben",
         "Sonstiges (bitte angeben)", "Bisher noch gar nicht"])

# ================= ABSCHNITT C =================
section("Abschnitt C – Selbsteinschätzung")

question(8, "Wie schätzen Sie Ihre eigenen Kompetenzen im Umgang mit KI ein?",
         "Matrix / Likert (Einfachauswahl je Zeile)", True)
matrix(
    ["Ich fühle mich im Umgang mit KI-Tools sicher.",
     "Ich kann gute, zielführende Eingaben (Prompts) formulieren.",
     "Ich kann KI-Ergebnisse kritisch einordnen und überprüfen.",
     "Ich kenne die Chancen und Grenzen von KI gut.",
     "Ich kenne grundlegende Datenschutzaspekte beim KI-Einsatz."],
    ["Trifft nicht zu", "Trifft eher nicht zu", "Teils/teils", "Trifft eher zu", "Trifft voll zu"]
)

# ================= ABSCHNITT D =================
section("Abschnitt D – Einstellung zur Einführung von KI an der Schule")

question(9, "Inwieweit stimmen Sie den folgenden Aussagen zu?",
         "Matrix / Likert (Einfachauswahl je Zeile)", True)
matrix(
    ["Ich stehe dem Einsatz von KI an unserer Schule grundsätzlich offen gegenüber.",
     "Ich erwarte einen konkreten Nutzen für meine tägliche Arbeit.",
     "Ich erhoffe mir eine spürbare Entlastung.",
     "Ich sehe Möglichkeiten, die Qualität meines Unterrichts zu verbessern.",
     "Ich habe Vorbehalte gegenüber dem Einsatz von KI.",
     "Ich wünsche mir klare schulische Regeln und Leitlinien zur KI-Nutzung."],
    ["Trifft nicht zu", "Trifft eher nicht zu", "Teils/teils", "Trifft eher zu", "Trifft voll zu"]
)

question(10, "Welche Hoffnungen oder Erwartungen verbinden Sie mit KI an unserer Schule? (optional)",
         "Freitext (lang)", False)

# ================= ABSCHNITT E =================
section("Abschnitt E – Mögliche Bedenken und Unterstützungsbedarf")
doc.add_paragraph(
    "Im Folgenden geht es um Aspekte, die beim Einsatz von KI eine Rolle spielen können. "
    "Bitte geben Sie an, wie bedeutsam Ihnen die jeweiligen Punkte erscheinen.",
    style="Intense Quote")

question(11, "Wie wichtig sind Ihnen die folgenden Aspekte bei der Einführung von KI?",
         "Matrix / Likert (Einfachauswahl je Zeile)", True)
matrix(
    ["Datenschutz und sichere Datenverarbeitung",
     "Rechtliche Klarheit und Sicherheit",
     "Vermeidung zusätzlicher Arbeitsbelastung",
     "Ausreichend Zeit zur Einarbeitung",
     "Verlässliche Qualität der KI-Ergebnisse",
     "Umgang mit möglicher Täuschung durch Schülerinnen und Schüler",
     "Auswirkungen auf Prüfungen und Leistungsbewertung",
     "Vermeidung zu starker Abhängigkeit von Technik",
     "Gleichmäßige, faire Nutzung im Kollegium"],
    ["Unwichtig", "Eher unwichtig", "Teils/teils", "Eher wichtig", "Sehr wichtig"]
)

question(12, "Wie wichtig wäre Ihnen eine klare schulische Leitlinie zur KI-Nutzung?",
         "Bewertungsskala 1–5", True)
doc.add_paragraph("1 = gar nicht wichtig … 5 = sehr wichtig", style="List Bullet")

question(13, "Welche Bedenken oder offenen Fragen möchten Sie ansprechen? (optional)",
         "Freitext (lang)", False)

# ================= ABSCHNITT F =================
section("Abschnitt F – Wünsche an eine schulische KI-Lösung")

question(14, "Welche Anforderungen sollte eine schulische KI-Lösung erfüllen? (Mehrfachauswahl)",
         "Mehrfachauswahl", True)
options(["Einfache, intuitive Bedienung", "Schulische Lizenz / einheitlicher Zugang",
         "Datenschutzkonforme Nutzung", "Vorab geprüfte, empfohlene Tools",
         "Vorlagen für Unterrichtsmaterial", "Unterstützung bei Arbeitsblättern",
         "Unterstützung bei Klassenarbeiten", "Unterstützung bei Differenzierung",
         "Unterstützung bei Korrektur und Feedback", "Unterstützung bei Verwaltungsaufgaben",
         "Zentrale Prompt-Bibliothek (Sammlung bewährter Eingaben)",
         "Fortbildungsangebote", "Austauschmöglichkeiten im Kollegium",
         "Sonstiges (bitte angeben)"])

question(15, "Welche drei dieser Punkte sind Ihnen am wichtigsten? (Bitte bis zu 3 auswählen)",
         "Mehrfachauswahl (max. 3) / alternativ Ranking", False)
options(["Einfache Bedienung", "Schulische Lizenz / einheitlicher Zugang",
         "Datenschutzkonforme Nutzung", "Geprüfte Tools", "Vorlagen & Materialunterstützung",
         "Korrektur und Feedback", "Verwaltungsaufgaben", "Prompt-Bibliothek",
         "Fortbildung", "Kollegialer Austausch"])

# ================= ABSCHNITT G =================
section("Abschnitt G – Fortbildungsbedarf")

question(16, "Zu welchen Themen wünschen Sie sich Fortbildung? (Mehrfachauswahl)",
         "Mehrfachauswahl", True)
options(["Grundlagen: Was ist KI und wie funktioniert sie?",
         "Praktische Nutzung gängiger Tools (ChatGPT, Claude, Gemini u. a.)",
         "Prompting für Lehrkräfte (gute Eingaben formulieren)",
         "KI für die Unterrichtsvorbereitung", "KI für Arbeitsblätter und Klassenarbeiten",
         "KI für Differenzierung", "KI und Prüfungen / Leistungsbewertung",
         "KI und Datenschutz", "KI im jeweiligen Fachunterricht",
         "Best-Practice-Beispiele aus dem Kollegium",
         "Sonstiges (bitte angeben)", "Ich habe derzeit keinen Fortbildungsbedarf"])

question(17, "Welche Fortbildungsformate würden Sie bevorzugen? (Mehrfachauswahl)",
         "Mehrfachauswahl", False)
options(["Kurze schulinterne Workshops", "Kollegiale Kurzformate (Mikro-Fortbildungen)",
         "Selbstlernmaterialien zum eigenen Tempo", "Ganztägige bzw. längere Fortbildungen",
         "Externe Fortbildungen", "Voneinander lernen im Tandem / kleinen Gruppen",
         "Sonstiges (bitte angeben)"])

question(18, "Wären Sie bereit, eigene Erfahrungen oder Beispiele im Kollegium zu teilen?",
         "Einfachauswahl", False)
options(["Ja, gerne", "Eventuell, je nach Thema", "Eher nicht", "Nein"])

# ================= ABSCHNITT H =================
section("Abschnitt H – Freiwillige Angaben und Abschluss")
doc.add_paragraph(
    "Die folgenden Angaben sind freiwillig. Sie helfen uns, Bedarfe besser einzuordnen "
    "(z. B. nach Fachbereich). Sie können alle Felder frei lassen.",
    style="Intense Quote")

question(19, "Welchem Fachbereich fühlen Sie sich überwiegend zugehörig? (optional)",
         "Einfachauswahl", False)
options(["Allgemeinbildende Fächer", "Berufsbezogene/technische Fächer",
         "Wirtschaft / Verwaltung", "Sozialpädagogik / Pflege / Gesundheit",
         "Sprachen", "Naturwissenschaften", "Sonstiger Bereich (bitte angeben)",
         "Keine Angabe"])

question(20, "In welcher Funktion sind Sie überwiegend tätig? (optional)",
         "Einfachauswahl", False)
options(["Lehrkraft", "Lehrkraft mit Funktionsstelle / erweiterte Aufgaben",
         "Schulleitung / erweiterte Schulleitung", "Referendariat / im Vorbereitungsdienst",
         "Keine Angabe"])

question(21, "Name (vollkommen freiwillig – nur, wenn Sie für Rückfragen ansprechbar sein möchten):",
         "Freitext (kurz)", False)

question(22, "Möchten Sie uns noch etwas mitteilen? (optional)",
         "Freitext (lang)", False)

doc.add_paragraph()
end = doc.add_paragraph()
re = end.add_run("Vielen Dank für Ihre Teilnahme und Ihre offene Rückmeldung!")
re.bold = True
re.font.color.rgb = ACCENT

doc.save("/home/user/familienapp/KI-Umfrage_Engelbert-Bohn-Schule.docx")
print("OK – docx erstellt")
